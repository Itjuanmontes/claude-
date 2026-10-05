#!/usr/bin/env python3
"""Invitation to Bid .docx generator — matches the user's approved reference
(invitation-to-bid-test-send-1.docx) exactly:
- header band #10233D, title #10233D, "Not specified" for missing values
- green #087443 footer band with help line + phone
- fuller "Click here ..." placeholders, "Click qty"/"Click unit" in bid form
- sections numbered 3, 4, 5 ($/SQF pricing), 6 to mirror the reference
- $/SQF unit-pricing section (trade-aware, price boxes NEVER prefilled)

Usage: invitation_docx.py <contacts.json> <output_dir>
"""
import json, re, sys, os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), '..'))
from solicitation_fields import MISSING, extract_fields

DARK = RGBColor(0x10, 0x23, 0x3D)
DARK_HEX = "10233D"
GREEN_HEX = "087443"
INK = RGBColor(0x1C, 0x19, 0x16)
GRAY = RGBColor(0x47, 0x54, 0x67)
PH = RGBColor(0x8A, 0x83, 0x76)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FILL_HEX = "FFF3C4"
FONT = "Calibri"

BRAND_UPPER = "LEAD MAGNET"
BRAND = "Lead Magnet"
PHONE = "855-799-8420"


def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:fill'), color)
    tcPr.append(shd)


def run(p, text, bold=False, size=11, color=INK, italic=False):
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    r.font.size = Pt(size)
    r.font.name = FONT
    if color:
        r.font.color.rgb = color
    return r


def cell_run(cell, text, bold=False, size=11, color=INK, italic=False, center=False):
    cell.text = ""
    p = cell.paragraphs[0]
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return run(p, text, bold=bold, size=size, color=color, italic=italic)


def fill_box(cell, placeholder="Click here to type"):
    cell_run(cell, placeholder, size=11, color=PH, italic=True, center=True)
    shade(cell, FILL_HEX)


def header_row(table, headers):
    for j, hh in enumerate(headers):
        cell_run(table.cell(0, j), hh, bold=True, size=11, color=WHITE)
        shade(table.cell(0, j), DARK_HEX)


def label_cell(cell, text):
    cell_run(cell, text, bold=True, size=11, color=INK)


def section(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    run(p, text, bold=True, size=13, color=INK)


def grid_table(doc, nrows, ncols, widths=None):
    t = doc.add_table(rows=nrows, cols=ncols)
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    if widths:
        for j, w in enumerate(widths):
            for i in range(nrows):
                t.cell(i, j).width = Inches(w)
    return t


SQF_MAP = [
    (('roof',), ["Tear-off & disposal", "Roof system installation", "Underlayment, flashing & penetrations"]),
    (('floor',), ["Flooring installation", "Subfloor preparation & leveling", "Transitions & trim"]),
    (('paint',), ["Interior painting", "Exterior painting", "Prep, prime & caulk"]),
    (('concrete', 'mason', 'brick', 'stone'), ["Concrete flatwork", "Masonry / block work", "Demo & surface prep"]),
    (('drywall',), ["Drywall hang", "Finish — tape, float & texture"]),
    (('sid',), ["Siding installation", "Trim, soffit & fascia"]),
    (('tile',), ["Tile installation", "Surface prep & waterproofing"]),
    (('carpentr', 'framing'), ["Rough framing", "Finish carpentry & trim"]),
    (('electric',), ["Electrical rough-in", "Electrical finish — devices & fixtures"]),
    (('plumb',), ["Plumbing rough-in", "Plumbing finish — fixtures & trim"]),
    (('hvac', 'heating', 'cooling', 'air conditioning'), ["HVAC rough-in / ductwork", "Equipment set, startup & commissioning"]),
    (('landscap',), ["Grading & prep", "Planting & sod"]),
    (('fence',), ["Fence installation", "Gates & hardware"]),
    (('window', 'door'), ["Window / door installation", "Trim & sealing"]),
    (('deck',), ["Deck framing & structure", "Decking & railings"]),
]


def sqf_items(trade, title):
    blob = f"{trade} {title}".lower()
    for keys, items in SQF_MAP:
        if any(k in blob for k in keys):
            return items
    return ["Base scope", "Additional scope"]


def extract_solicitation(bid):
    return extract_fields(bid)['Solicitation']


def build_invitation(contact, bid):
    doc = Document()
    doc.styles['Normal'].font.name = FONT
    doc.styles['Normal'].font.size = Pt(11)
    for s in doc.sections:
        s.top_margin = Inches(0.6)
        s.bottom_margin = Inches(0.6)
        s.left_margin = Inches(0.7)
        s.right_margin = Inches(0.7)

    f = extract_fields(bid)
    solicitation = f['Solicitation']
    sam = bid.get('sam') or {}
    trade = contact.get('category') or MISSING

    # Header band
    hb = grid_table(doc, 1, 1)
    c0 = hb.cell(0, 0)
    shade(c0, DARK_HEX)
    c0.text = ""
    p = c0.paragraphs[0]
    run(p, BRAND_UPPER, bold=True, size=14, color=WHITE)
    p.add_run("\n")
    run(p, f"Invitation to Bid · {solicitation} · fill this file in and email it back",
        size=10, color=WHITE)

    p = doc.add_paragraph()
    run(p, "Invitation to Bid", bold=True, size=22, color=DARK)
    p = doc.add_paragraph()
    run(p, solicitation, bold=True, size=12, color=INK)
    p = doc.add_paragraph()
    run(p, "Type in the yellow boxes, save this Word file, and email it back to the person who sent it. "
           "Do not bid from the email body. Bid the owner package and the latest addendum.",
        size=11, color=INK)

    # Job details
    section(doc, "Job details")
    poc_line = " · ".join(x for x in f['POC'].values() if x) or MISSING
    issued_for = contact.get('business_name', MISSING)
    city, state = contact.get('city') or "", contact.get('state') or ""
    loc = ", ".join(x for x in [city, state] if x)
    if loc:
        issued_for += f" · {loc}"
    jd = [
        ("Issued by", BRAND),
        ("Issued for", issued_for),
        ("Trade", trade),
        ("Owner / Agency", f['Agency']),
        # Place of performance only. Never the contractor's own state.
        ("Place", f['Place']),
        ("Solicitation", solicitation),
        ("Package", bid.get('type_label') or "Open Opportunity"),
        ("Documents issued", MISSING),
        ("Posted", f['Posted']),
        ("Bids due", f['Deadline']),
        ("NAICS", f['NAICS']),
        ("Set-aside", f['Set-aside']),
        ("Official POC", poc_line),
        ("Official listing", bid.get('url') or MISSING),
    ]
    t = grid_table(doc, len(jd) + 1, 2, widths=[1.7, 5.3])
    header_row(t, ("Field", "Value"))
    for i, (lab, v) in enumerate(jd, start=1):
        label_cell(t.cell(i, 0), lab)
        cell_run(t.cell(i, 1), v)
    p = doc.add_paragraph()
    run(p, "This is a working invitation so you can price the job. It is not a substitute for the owner solicitation. "
           f"Official award is by the owner, not {BRAND}.", size=10, color=GRAY)

    # 1. Bidder information
    section(doc, "1. Bidder information  —  fill this in")
    fields = [
        ("Legal company name", "Click here to type company name"),
        ("DBA / trade name", "Click here if different"),
        ("Contact name", "Click here to type your name"),
        ("Title", "Click here — estimator / owner"),
        ("Phone", "Click here — cell / office"),
        ("Email", "Click here — email we should reply to"),
        ("Address / city / state / zip", "Click here to type mailing address"),
        ("License # / state", "Click here to type license"),
        ("FEIN (optional)", "Click here to type FEIN"),
        ("Years in this trade", "Click here to type years"),
    ]
    t = grid_table(doc, len(fields) + 1, 2, widths=[2.3, 4.7])
    header_row(t, ("Item", "Your answer"))
    for i, (lab, ph) in enumerate(fields, start=1):
        label_cell(t.cell(i, 0), lab)
        fill_box(t.cell(i, 1), ph)

    # 2. Scope
    section(doc, "2. Scope (read this)")
    desc = sam.get('description') or bid.get('title', '')
    scope_first = ""
    for sent in [s.strip(" •") for s in re.split(r'(?<=[.!?])\s+', desc) if s.strip()][:1]:
        scope_first = sent.strip()
    if scope_first:
        doc.add_paragraph(f"•  {scope_first} — scope per the owner solicitation and amendments ({trade}).")
    for line in ["•  Furnish labor, equipment, materials, haul-off, and protection unless the owner package states otherwise.",
                 "•  Daily cleanup. Do not leave debris on site.",
                 "•  Do not bid from this document body alone. Pull drawings and addenda from the official listing."]:
        doc.add_paragraph(line)

    # 3. Bid form
    section(doc, "3. Bid form  —  fill unit price and total")
    trade_work = f"{trade} work per the owner scope" if trade != MISSING else "Trade work per the owner scope"
    t = grid_table(doc, 6, 6, widths=[0.6, 2.6, 0.6, 0.6, 1.3, 1.3])
    header_row(t, ("#", "Description", "Qty", "Unit", "Unit price $", "Total $"))
    for i, (num, dsc, qty, unit) in enumerate([
        ("1", "Mobilization / protection", "1", "LS"),
        ("2", trade_work, "1", "LS"),
        ("3", "Haul-off / dump fees", "1", "LS"),
    ], start=1):
        cell_run(t.cell(i, 0), num, center=True)
        cell_run(t.cell(i, 1), dsc)
        cell_run(t.cell(i, 2), qty, center=True)
        cell_run(t.cell(i, 3), unit, center=True)
        fill_box(t.cell(i, 4), "Click here $")
        fill_box(t.cell(i, 5), "Click here $")
    cell_run(t.cell(4, 0), "Other", center=True)
    cell_run(t.cell(4, 1), "Other (describe)")
    fill_box(t.cell(4, 2), "Click qty")
    fill_box(t.cell(4, 3), "Click unit")
    fill_box(t.cell(4, 4), "Click here $")
    fill_box(t.cell(4, 5), "Click here $")
    label_cell(t.cell(5, 1), "BASE BID TOTAL")
    fill_box(t.cell(5, 5), "Click here — total $")

    # 4. Alternates
    section(doc, "4. Alternates / notes  —  fill if needed")
    alt = [
        ("Alternate 1 (describe + $)", "Click here — alt 1"),
        ("Alternate 2 (describe + $)", "Click here — alt 2"),
        ("Exclusions", "Click here — what is not included"),
        ("Lead time / start", "Click here — days after award"),
        ("Questions / clarifications", "Click here — anything we should know"),
    ]
    t = grid_table(doc, len(alt) + 1, 2, widths=[2.3, 4.7])
    header_row(t, ("Item", "Your answer"))
    for i, (lab, ph) in enumerate(alt, start=1):
        label_cell(t.cell(i, 0), lab)
        fill_box(t.cell(i, 1), ph)

    # 5. Unit pricing $/SQF
    section(doc, "5. Unit pricing  —  price per square foot")
    p = doc.add_paragraph()
    run(p, "Enter your price per square foot for each line below. Leave blank any line that does not apply.",
        size=10, color=GRAY, italic=True)
    items = sqf_items(trade, bid.get('title') or '')
    t = grid_table(doc, len(items) + 2, 3, widths=[3.2, 1.1, 2.7])
    header_row(t, ("Work item", "Unit", "Your price"))
    for i, it in enumerate(items, start=1):
        label_cell(t.cell(i, 0), it)
        cell_run(t.cell(i, 1), "$ / SQF", center=True, color=GRAY)
        fill_box(t.cell(i, 2), "Click here $ / SQF")  # NEVER prefilled
    n = len(items) + 1
    label_cell(t.cell(n, 0), "Notes / inclusions")
    t.cell(n, 1).merge(t.cell(n, 2))
    fill_box(t.cell(n, 1), "Click here to type inclusions, exclusions, assumptions")

    # 6. Sign and return
    section(doc, "6. Sign and return")
    sig = [
        ("Authorized signature (type full name)", "Click here to type signature"),
        ("Printed name", "Click here to type printed name"),
        ("Title", "Click here to type title"),
        ("Date", "Click here to type date"),
    ]
    t = grid_table(doc, len(sig) + 1, 2, widths=[2.3, 4.7])
    header_row(t, ("Item", "Your answer"))
    for i, (lab, ph) in enumerate(sig, start=1):
        label_cell(t.cell(i, 0), lab)
        fill_box(t.cell(i, 1), ph)

    # Footer band
    doc.add_paragraph()
    fb = grid_table(doc, 1, 1)
    fc = fb.cell(0, 0)
    shade(fc, GREEN_HEX)
    fc.text = ""
    p = fc.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run(p, "If you'd like help filling this invitation out or have questions, call us.", size=11, color=WHITE)
    p.add_run("\n")
    run(p, PHONE, bold=True, size=16, color=WHITE)
    p.add_run("\n")
    run(p, "Save this Word file and email it back to the person who sent it.", size=10, color=WHITE)
    p = doc.add_paragraph()
    run(p, f"Working invitation from {BRAND}. Official award is by the owner, not {BRAND}. Bid the owner solicitation on SAM.gov.",
        size=9, color=GRAY, italic=True)

    return doc


def main():
    src, outdir = sys.argv[1], sys.argv[2]
    os.makedirs(outdir, exist_ok=True)
    contacts = json.load(open(src))
    n = 0
    for c in contacts:
        spid = re.sub(r'\W+', '_', str(c.get('spid') or c.get('business_name', 'contact')))[:40]
        for b in c['bids']:
            if b.get('expired'):
                continue
            doc = build_invitation(c, b)
            title_slug = re.sub(r'\W+', '_', (b.get('title') or 'bid'))[:40].strip('_')
            fn = f"{spid}_bid{b.get('slot', n)}_{title_slug}_invitation-to-bid.docx"
            doc.save(os.path.join(outdir, fn))
            n += 1
    print(f"generated {n} invitations in {outdir}")


if __name__ == '__main__':
    main()
